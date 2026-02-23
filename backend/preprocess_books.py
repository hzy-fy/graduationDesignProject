import os
import re
import pdfplumber
from tqdm import tqdm
import pytesseract
from pdf2image import convert_from_path
import docx

# 定位数据目录
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'data')
BOOKS_DIR = os.path.join(DATA_DIR, 'booksdata')
PROCESSED_DIR = os.path.join(DATA_DIR, 'processed')

# OCR 引擎配置 (如果已添加到环境变量 PATH，则无需配置)
# pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

def clean_text(text):
    """
    对提取的文本进行清洗
    """
    if not text:
        return ""
        
    lines = text.split('\n')
    cleaned_lines = []
    for line in lines:
        line = line.strip()
        # 跳过纯数字行 (页码)
        if re.match(r'^\d+$', line):
            continue
        # 跳过过短的行
        if len(line) < 2:
            continue
        cleaned_lines.append(line)
    
    text = '\n'.join(cleaned_lines)

    # 统一术语
    text = text.replace('寒底体质', '寒性体质')
    text = text.replace('热底体质', '热性体质')
    text = text.replace('寒底', '寒性')
    text = text.replace('热底', '热性')

    # 去除乱码
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
    
    # 去除多余空白
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'\n\s*\n', '\n', text)

    return text.strip()

def process_with_ocr(pdf_path):
    """
    使用 OCR 处理扫描件 PDF
    """
    print(f"  Starting OCR for {os.path.basename(pdf_path)}... (This may take a while)")
    text_content = []
    
    try:
        # 将 PDF 转为图片 (dpi=200 平衡速度与精度)
        # 如果报错 "Poppler not found"，请安装 Poppler 并添加 bin 目录到 PATH
        images = convert_from_path(pdf_path, dpi=200)
        
        # 跳过前几页
        start_page = 0
        if len(images) > 10:
            start_page = 5
            
        print(f"  Total pages (OCR): {len(images)}. Processing from page {start_page} (Testing first 10 pages)...")

        # 仅处理前 10 页用于测试
        test_limit = 10
        for i, image in enumerate(tqdm(images[start_page:start_page+test_limit])):
            # OCR 识别 (需安装 Tesseract 并配置 chi_sim 语言包)
            # lang='chi_sim+eng' 表示同时识别简体中文和英文
            try:
                text = pytesseract.image_to_string(image, lang='chi_sim+eng')
                if text:
                    cleaned = clean_text(text)
                    if cleaned:
                        text_content.append(cleaned)
            except pytesseract.TesseractError as e:
                print(f"    Tesseract Error on page {i}: {e}")
                continue
                    
    except Exception as e:
        print(f"  OCR Failed: {e}")
        print("  Tip: Please ensure Poppler is installed and added to PATH.")
        return None
        
    return "\n\n".join(text_content)

def process_pdf(pdf_path):
    print(f"Processing: {os.path.basename(pdf_path)}")
    text_content = []
    
    # 1. 尝试使用 pdfplumber 提取文本
    try:
        with pdfplumber.open(pdf_path) as pdf:
            start_page = 0
            if len(pdf.pages) > 10:
                start_page = 5
                
            # 只尝试提取前 10 页来判断是否为扫描件
            sample_pages = pdf.pages[start_page:start_page+10]
            sample_text = ""
            for page in sample_pages:
                sample_text += page.extract_text() or ""
            
            if len(sample_text.strip()) < 100:
                print("  Text extraction yielded very little content. Switching to OCR...")
                return process_with_ocr(pdf_path)
            
            # 如果是文本 PDF，继续提取全部
            print(f"  Text PDF detected. Extracting {len(pdf.pages)} pages...")
            for page in tqdm(pdf.pages[start_page:]):
                text = page.extract_text()
                if text:
                    cleaned = clean_text(text)
                    if cleaned:
                        text_content.append(cleaned)
                        
    except Exception as e:
        print(f"Error reading PDF {pdf_path}: {e}")
        return None
    
    return "\n\n".join(text_content)

def process_docx(docx_path):
    print(f"Processing DOCX: {os.path.basename(docx_path)}")
    text_content = []
    
    try:
        doc = docx.Document(docx_path)
        
        # doc.paragraphs 是正文段落列表，本身不包含页眉页脚
        # 页眉页脚需要通过 doc.sections[0].header.paragraphs 获取，我们不需要
        
        print(f"  Extracting text from {len(doc.paragraphs)} paragraphs...")
        
        for para in tqdm(doc.paragraphs):
            text = para.text.strip()
            if text:
                cleaned = clean_text(text)
                if cleaned:
                    text_content.append(cleaned)
                    
    except Exception as e:
        print(f"Error reading DOCX {docx_path}: {e}")
        return None
        
    return "\n".join(text_content)

def main():
    if not os.path.exists(PROCESSED_DIR):
        os.makedirs(PROCESSED_DIR)
        
    if not os.path.exists(BOOKS_DIR):
        print(f"Error: Books directory not found at {BOOKS_DIR}")
        return

    # 扫描 PDF 和 DOCX 文件
    files = [f for f in os.listdir(BOOKS_DIR) if f.lower().endswith(('.pdf', '.docx'))]
    print(f"Found {len(files)} supported files.")

    for filename in files:
        file_path = os.path.join(BOOKS_DIR, filename)
        txt_filename = os.path.splitext(filename)[0] + '.txt'
        txt_path = os.path.join(PROCESSED_DIR, txt_filename)
        
        if os.path.exists(txt_path):
            # 检查文件大小
            if os.path.getsize(txt_path) > 1024:
                print(f"Skipping {filename} (already processed)")
                continue
            
        content = None
        if filename.lower().endswith('.pdf'):
            content = process_pdf(file_path)
        elif filename.lower().endswith('.docx'):
            content = process_docx(file_path)
        
        if content and len(content.strip()) > 100:
            with open(txt_path, 'w', encoding='utf-8') as f:
                f.write(content)
            print(f"Saved processed text to {txt_path} (Size: {len(content)} chars)")
        else:
            print(f"Warning: No valid text extracted from {filename}.")

if __name__ == "__main__":
    main()
